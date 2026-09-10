"""Rebuild the W2001/interval manuscript figures and verify their source tables.

Run with a Python environment containing matplotlib, Pillow, openpyxl and PyMuPDF.
Use --skip-r only after a successful run of the primary R script.
"""
import argparse
import csv
from datetime import datetime
import gzip
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys

import pymupdf as fitz
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont, ImageStat

ANALYSIS = Path(__file__).resolve().parent
PROJECT = ANALYSIS.parents[2]
R_SCRIPT = ANALYSIS / "01_promoter_enhancer_interaction_resubmit_loop_annotation.R"
WORKBOOK = PROJECT / "outputs/promoter-window-anchor-sensitivity/promoter_window_anchor_sensitivity_atac50bp_pooled_5k_10k_25k_gene_lists_with_loop_count_tabs.xlsx"

# Manuscript numbering, not the older figure*_revised_* export numbering.
FIGURES = [
    ("1", "sequencing_and_loop_depth", "10 libraries; sample-level calls <2 Mb", "Read composition and sequencing-depth association; not causation."),
    ("2", "loop_counts_per_chr", "31,021 pooled exact calls <2 Mb", "Counts by chromosome and detection resolution; not restricted to regulatory calls."),
    ("3", "shared_loops", "45 library pairs per resolution; 10 libraries", "a: pairwise exact overlap mean +/- SD, including zero pairs. b: support in one versus multiple libraries."),
    ("4", "ctcf_and_gene_density", "22 rn7 chromosomes", "Predicted motif intervals in non-overlapping 1-Mb bins; Ensembl 113 genes, not the legacy NCBI gene catalog."),
    ("5", "density_plot", "30,932 pooled calls with full [-1,2] windows", "Motif midpoints, unique strand-aware Ensembl TSS sites and EPD promoter midpoints; KDE by resolution."),
    ("6", "ctcf_anchor_motif_counts", "31,021 calls / 62,042 anchors", "Descriptive motif counts at full anchor intervals; log2(count+1) retains zeros. No CTCF-count selection."),
    ("7", "flowchart", "59,000 sample calls -> 31,021 pooled -> 13,376 putative", "Counts read from this R run; W2001/interval/ATAC50; dual-promoter support requires a completed direction."),
    ("8", "circos_putative_regulatory_loops", "13,376 putative calls", "All chromosomes and chromosome 1; includes 2,907 directionally supported dual-promoter calls."),
    ("S1", "loop_qc_correlation_heatmap", "10 libraries; same <2-Mb counts as Fig. 1", "Pearson correlations; QC metrics are correlated with each other and do not establish causality."),
    ("S3", "pairwise_common_loop_percentage_by_resolution", "10 x 10 libraries at each resolution; calls <2 Mb", "Each cell is 100 x intersection / calls in the X-axis library. The matrix is not generally symmetric."),
    ("S4", "density_histograms", "Same 30,932 calls as Fig. 5", "200-bin feature-loop overlap histograms; repeated feature appearances in different loops are retained."),
    ("S5", "ctcf_histograms_by_chromosome", "Same 30,932 calls as Fig. 5", "Chromosome-wise predicted CTCF motif midpoint histograms, not experimentally measured binding."),
    ("S6", "tss_histograms_by_chromosome", "Same 30,932 calls as Fig. 5", "Chromosome-wise strand-aware Ensembl TSS histograms."),
    ("S7", "promoter_histograms_by_chromosome", "Same 30,932 calls as Fig. 5", "Chromosome-wise EPD promoter midpoint histograms."),
    ("S9", "circos_by_chromosome", "13,376 putative calls, 22 chromosome panels", "Per-chromosome view of exactly the Fig. 8 call set; display midpoints do not change interval-based selection."),
]


def read_tsv(path):
    opener = gzip.open if path.suffix == ".gz" else open
    with opener(path, "rt") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def require(condition, message):
    if not condition:
        raise ValueError(message)


def verify_and_index(results):
    source = results / "figure_source_data"
    metrics = {r["metric"]: r["value"] for r in read_tsv(source / "run_metrics.tsv")}
    require(read_tsv(source / "r_generation_complete.tsv")[0]["run_started"] == metrics["run_started"], "R run incomplete")
    expected = {"window_bp": "2001", "anchor_mode": "interval", "atac_min_bp": "50", "putative": "13376",
                "genes": "8775", "gene_loop_pairs": "19292", "pooled_lt2mb": "31021"}
    require(all(metrics[k] == v for k, v in expected.items()), "Unexpected production parameters or totals")

    genes = read_tsv(source / "putative_gene_counts.tsv")
    gene_counts = {r["ensembl_gene_id"]: int(r["n"]) for r in genes}
    require(len(genes) == len(gene_counts) == 8775, "Gene IDs missing or duplicated")
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    require(workbook.sheetnames[0] == "W2001_Interval_13376", "Unexpected first workbook sheet")
    rows = list(workbook.worksheets[0].iter_rows(values_only=True))
    require(rows[0] == ("ensembl_gene_id", "gene_symbol", "n"), "Workbook columns changed")
    excel_counts = {r[0]: int(r[2]) for r in rows[1:] if r[0]}
    require(len(rows) - 1 == len(excel_counts) and gene_counts == excel_counts, "Gene counts differ from workbook")
    excel_symbols = {r[0]: r[1] or "" for r in rows[1:] if r[0]}
    symbol_differences = [r["ensembl_gene_id"] for r in genes if (r["gene_symbol"] or "") != excel_symbols[r["ensembl_gene_id"]]]
    require(not symbol_differences, f"Gene symbols differ from workbook: {symbol_differences[:10]}")
    workbook.close()

    loops = read_tsv(source / "loop_evidence.tsv.gz")
    require(len(loops) == len({r["loop_id"] for r in loops}) == 31021, "Loop grain mismatch")
    existing_resource_file = results / "revised_pooled_loop_annotation_resource.tsv"
    if not existing_resource_file.exists():
        existing_resource_file = ANALYSIS / "results/revised_pooled_loop_annotation_resource.tsv"
    existing_resource = read_tsv(existing_resource_file)
    existing_by_id = {r["loop_id"]: r for r in existing_resource}
    comparison_columns = ["chr1", "start1", "end1", "chr2", "start2", "end2", "resolution",
                          "n_direct_anchor_sides", "revised_putative_regulatory_support", "revised_major_category",
                          "revised_detailed_category", "supported_regulatory_direction"]
    require(set(existing_by_id) == {r["loop_id"] for r in loops}, "Production pooled loop IDs changed")
    require(all(r[k] == existing_by_id[r["loop_id"]][k] for r in loops for k in comparison_columns),
            "Production coordinates or regulatory classification changed")
    putative_ids = {r["loop_id"] for r in loops if r["revised_putative_regulatory_support"] == "TRUE"}
    assignments = read_tsv(source / "putative_gene_assignments.tsv.gz")
    require(len(putative_ids) == 13376 and {r["loop_id"] for r in assignments} == putative_ids, "Gene assignment loop coverage mismatch")
    require(len({(r["loop_id"], r["gene_id"]) for r in assignments}) == len(assignments) == 19292, "Gene-loop pair mismatch")
    require(sum(n >= 2 for n in gene_counts.values()) == 4631, "Table S3 subset count mismatch")
    circos = read_tsv(source / "circos_counts.tsv")
    require(sum(int(r["n_loops"]) for r in circos) == 13376, "Circos denominator mismatch")
    resolution_counts = {res: sum(int(r["n_loops"]) for r in circos if r["resolution"] == res) for res in ("5K", "10K", "25K")}
    require(resolution_counts == {"5K": 1896, "10K": 4716, "25K": 6764}, "Circos resolution counts mismatch")
    pairs = read_tsv(source / "figure3_pairwise_counts.tsv")
    require(len(pairs) == 135 and all(sum(r["resolution"] == res for r in pairs) == 45 for res in resolution_counts), "Pairwise figure must use 45 pairs per resolution")
    for res in resolution_counts:
        matrix = read_tsv(source / f"figureS3_{res}.tsv")
        require(len(matrix) == 100 and all(0 <= float(r["value"]) <= 100 for r in matrix), "Pairwise matrix invalid")
        require(all(float(r["value"]) == 100 for r in matrix if r["X_strain"] == r["Y_strain"]), "Pairwise diagonal not 100%")
    require(len(read_tsv(source / "positional_loop_universe.tsv.gz")) == 30932, "Positional universe mismatch")
    motif_counts = read_tsv(source / "figure6_motif_counts.tsv.gz")
    require(len(motif_counts) == 62042, "CTCF histogram must retain all anchors")
    bins = read_tsv(source / "figure4_motif_bins.tsv")
    previous_end = {}
    for row in bins:
        start, end = int(row["Start"]), int(row["End"])
        require(start == previous_end.get(row["Chr"], 0) and 0 < end - start <= 1000000,
                "Figure 4 bins overlap, have gaps, or have invalid width")
        previous_end[row["Chr"]] = end

    manifest = []
    preview = Image.new("RGB", (1800, 2100), "white")
    draw = ImageDraw.Draw(preview)
    font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 22)
    for i, (number, stem, population, note) in enumerate(FIGURES):
        base = "revision_figure" + number + "_" + stem
        png, pdf = results / (base + ".png"), results / (base + ".pdf")
        require(png.is_file() and pdf.is_file(), f"Missing Figure {number}")
        run_time = datetime.strptime(metrics["run_started"], "%Y-%m-%dT%H:%M:%S%z").timestamp()
        require(min(png.stat().st_mtime, pdf.stat().st_mtime) >= run_time, f"Stale Figure {number}")
        with Image.open(png) as original:
            img = original.convert("RGB")
            width, height = img.size
            require(min(width, height) >= 900, f"Figure {number} resolution too low")
            require(max(ImageStat.Stat(img).stddev) > 5, f"Figure {number} appears blank")
            thumb = img.copy()
            thumb.thumbnail((570, 370))
        with fitz.open(pdf) as document:
            require(len(document) == 1, f"Figure {number} PDF page count invalid")
            pix = document[0].get_pixmap(matrix=fitz.Matrix(0.5, 0.5), alpha=False)
            require(max(ImageStat.Stat(Image.frombytes("RGB", (pix.width, pix.height), pix.samples)).stddev) > 5,
                    f"Figure {number} PDF appears blank")
        x, y = (i % 3) * 600, (i // 3) * 420
        draw.text((x + 15, y + 10), f"Figure {number}", fill="black", font=font)
        preview.paste(thumb, (x + (600 - thumb.width) // 2, y + 45 + (370 - thumb.height) // 2))
        manifest.append(dict(figure=number, png=str(png), pdf=str(pdf), population=population,
                             width_px=width, height_px=height, png_sha256=sha256(png), pdf_sha256=sha256(pdf), note=note))
    preview.save(results / "manuscript_figure_index.png")
    with (results / "manuscript_figure_manifest.tsv").open("w") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(manifest[0]), delimiter="\t")
        writer.writeheader()
        writer.writerows(manifest)
    checks = dict(status="passed", run=metrics, workbook=str(WORKBOOK), workbook_sha256=sha256(WORKBOOK),
                  workbook_gene_ids_symbols_counts_equal=True, resolution_counts=resolution_counts,
                  pooled_coordinates_and_classifications_unchanged=True,
                  zero_motif_anchors=sum(int(r["motif_count"]) == 0 for r in motif_counts),
                  both_zero_motif_loops=sum(r["predicted_ctcf_motif_interval_count_anchor1"] == "0" and
                                            r["predicted_ctcf_motif_interval_count_anchor2"] == "0" for r in loops),
                  validated_figure_pairs=len(manifest),
                  source_sha256={str(p): sha256(p) for p in (R_SCRIPT, ANALYSIS / "draw_flowchart_revision.py", Path(__file__), PROJECT / "funcs_enhancer.R")})
    (source / "validation.json").write_text(json.dumps(checks, indent=2) + "\n")
    write_report(results, manifest, checks)
    print(f"PASS: {len(manifest)} PNG/PDF pairs; 13,376 loops, 8,775 genes and 19,292 gene-loop pairs match W2001_Interval_13376.", flush=True)


def write_report(results, manifest, checks):
    lines = ["# Manuscript Figure Replacement Audit", "", f"R run: {checks['run']['run_started']}", "",
        "## Result", "",
        "All 15 requested figure counterparts exist as regenerated PNG/PDF pairs. Numerical checks passed; see figure_source_data/validation.json. The files below use the manuscript numbering, not the older figure*_revised_* export numbering.", "",
        "- Locked settings: rn7; Ensembl/EPD TSS +/-1000 bp (2001 bp); any >=1-bp overlap with the full 5/10/25-kb anchor interval. ATAC support threshold is >=50 bp, not the promoter-overlap threshold.",
        "- ATAC union is split after removing Ensembl/EPD TSS +/-1000-bp windows. Single-promoter calls require opposite-anchor non-TSS ATAC; dual-promoter calls require promoter-window ATAC AND opposite non-TSS ATAC completing at least one direction.",
        "- No nearest-TSS 200-kb restriction, gene-inside-loop condition, or CTCF-count filter is used to select putative calls.",
        "- 59,000 sample-level calls -> 31,778 exact coordinate calls -> 31,021 calls <2 Mb -> 13,376 putative calls (5K: 1,896; 10K: 4,716; 25K: 6,764). Exact calls at different resolutions are not collapsed to unique biological interactions.",
        "- Workbook W2001_Interval_13376 matches all 8,775 gene IDs/symbols/counts and 19,292 gene-loop pairs. Genes with n>=2: 4,631 (Table S3).", "",
        "## Replacement Files", "", "| Figure | PNG / PDF | Population and interpretation |", "|---|---|---|"]
    for row in manifest:
        lines.append(f"| {row['figure']} | [PNG]({row['png']}) / [PDF]({row['pdf']}) | {row['population']}. {row['note']} |")
    lines += ["", "## Corrections Made", "",
        "- Restored the requested Fig. 6 motif-count histogram; zero counts retained via log2(count+1). The existing TSS/promoter-status stacked bar is retained as an additional figure, not mistaken for this replacement.",
        "- Fig. 3a now computes the manuscript's pairwise mean/SD across all 45 pairs at each resolution, rather than per-library multi-library-supported counts. Fig. 3b depth labels now come from the actual QC input, not hand-entered values.",
        "- Fig. 1 displays small P values in scientific notation rather than P=0.0000. S1 and S3 consistently apply the <2-Mb call universe; S3 explicitly identifies its X-axis denominator.",
        "- Fig. 4 genomic counting bins now correctly convert 0-based display intervals to non-overlapping 1-based GRanges bins. Motif intervals spanning a bin boundary can legitimately overlap two distinct bins.",
        "- Restored the clipped Fig. 4 panel-a label and explicitly exported Fig. 4/8 on white backgrounds.",
        "- Fig. 7 reads current R-run counts, names 10 libraries (not 10 independent strains), and includes the missing promoter-accessibility condition for dual-promoter calls.",
        "- Figure output paths now honor RESUBMIT_OUTPUT_DIR; non-default sensitivity settings cannot silently overwrite manuscript figures.", "",
        "## Essential Caption Caveats", "",
        "- Figures 1-6/S1/S3-S7 are QC, genome-wide or preselection context, not all restricted to 13,376 loops. Restricting those positional plots to TSS-selected loops would make TSS enrichment partly circular.",
        "- Fig. 5/S4-S7 exclude 89 calls whose expanded [-1,2] windows cross chromosome boundaries: 30,932 remain. Positions 0 and 1 denote anchor midpoints; this display normalization does not reinstate midpoint-based gene assignment. Features can appear in several loops; these counts are not counts of independent features.",
        "- CTCF panels show predicted motif intervals, not ChIP-validated binding, occupancy, or a demonstrated threshold separating true loops from noise. The legacy caption claiming low-count calls are noise must not be reused.",
        "- Fig. 6 UP/DOWN means the lower-/higher-coordinate anchor, not transcriptional upstream/downstream for every gene. Zero anchors are shown at log2(0+1)=0.",
        "- The 50-bp test requires overlap with at least one contiguous reduced ATAC interval (or non-TSS residual interval), not a sum of disconnected short overlaps. For dual-promoter calls, ATAC is tested against a full assigned promoter window, not only its intersection with the anchor; gene assignment is filtered by supported anchor direction, not individually by each gene's promoter accessibility.",
        "- RNA promoter/TSS annotations and open chromatin support candidate regulation, not causal enhancer activity. The ATAC data are not matched measurements from the 10 Hi-C libraries. No wet-lab validation is implied.",
        "- Circos links use anchor midpoints for display. The current drawing code does not explicitly use the legacy caption's logarithmic arc-height scaling; omit that claim.",
        "- Fig. 4 and Circos composite PDFs contain raster panels; PNGs and PDFs were checked for nonblank rendering, but these are not wholly vector figures.",
        f"- Current motif counts: {checks['zero_motif_anchors']:,} zero-count anchors; {checks['both_zero_motif_loops']:,} calls with zeros at both anchors. Do not retain the old '111 of 31,019' caption.", "",
        "## Sources and Reproduction", "",
        f"- Main R source: [{R_SCRIPT.name}]({R_SCRIPT})",
        f"- Flowchart: [draw_flowchart_revision.py]({ANALYSIS / 'draw_flowchart_revision.py'})",
        f"- Rebuild and verification: [rebuild_manuscript_figures.py]({Path(__file__).resolve()})",
        f"- Normalized input cache: `{ANALYSIS / 'cache_data'}`; raw-input provenance: [input_files.tsv]({results / 'figure_source_data/input_files.tsv'}).",
        "- This run recomputed annotations and plots from the existing normalized coordinate cache; it did not realign Hi-C reads or redo liftover. Production pooled coordinates/classifications were independently compared and are unchanged.",
        f"- Workbook: [{WORKBOOK.name}]({WORKBOOK})",
        f"- Exact figure mapping and hashes: [manuscript_figure_manifest.tsv]({results / 'manuscript_figure_manifest.tsv'}). Per-figure source tables and checks: `{results / 'figure_source_data'}`.",
        f"- Quick visual index: [manuscript_figure_index.png]({results / 'manuscript_figure_index.png'}).", "",
        "```bash", f"cd {PROJECT}",
        "uv run --with matplotlib --with pillow --with openpyxl --with pymupdf python r_files/revision/revision_main/rebuild_manuscript_figures.py", "```", "",
        "Do not use the older figure*_revised_* files or canonical_JASPAR sensitivity images as these 15 replacements. No old files were deleted."
    ]
    (results / "MANUSCRIPT_FIGURE_REPLACEMENT_REPORT.md").write_text("\n".join(lines) + "\n")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--skip-r", action="store_true")
    parser.add_argument("--output-dir", type=Path, default=ANALYSIS / "results" / "2nd_resubmission")
    args = parser.parse_args()
    results = args.output_dir.resolve()
    results.mkdir(parents=True, exist_ok=True)
    env = os.environ.copy()
    env.update(ENHANCER_PROJECT_DIR=str(PROJECT), RESUBMIT_ANALYSIS_DIR=str(ANALYSIS),
               RESUBMIT_OUTPUT_DIR=str(results), PROMOTER_WINDOW_FLANK_BP="1000", PROMOTER_ANCHOR_MATCH_MODE="interval",
               ATAC_MINIMUM_OVERLAP_BP="50", PROMOTER_WINDOW_SENSITIVITY_ONLY="0", PROMOTER_WINDOW_GENE_LIST_ONLY="0")
    if not args.skip_r:
        print("Recomputing annotations and R figures; progress log: " + str(results / "manuscript_figure_generation.log"), flush=True)
        with (results / "manuscript_figure_generation.log").open("w") as handle:
            subprocess.run([shutil.which("Rscript") or "/usr/local/bin/Rscript", str(R_SCRIPT)], cwd=PROJECT,
                           env=env, stdout=handle, stderr=subprocess.STDOUT, check=True)
    print("Regenerating flowchart from R-run metrics", flush=True)
    subprocess.run([sys.executable, str(ANALYSIS / "draw_flowchart_revision.py")], env=env, check=True)
    verify_and_index(results)


if __name__ == "__main__":
    main()
